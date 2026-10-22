#!/usr/bin/env python3
"""Bounded streaming QC of raw ZIP/tar representations. No raw files altered.

Artifact types, official split membership and unknown schemas remain explicit.
Competition solution files never become approved/publicly usable labels merely
because they occur in an upstream bundle. No training or threshold optimization.
"""
import argparse,collections,csv,hashlib,io,json,math,re,sqlite3,stat,tarfile,zipfile
from pathlib import Path
from rnagym import ROOT,StateDB,get_target,immutable_json,read_manifest,safe_member,sha256
from qc_sequence_files import Numeric

MAX_EXPANDED=32*1024**3
MAX_LINE=16*1024**2


def schema_kind(fields):
 if {'position_id','sequence_id','sequence'}.issubset(fields):return 'long_nucleotide_probing'
 if 'reactivity_0001' in fields:return 'wide_nucleotide_probing'
 if {'sequence','DMS_score'}.issubset(fields):return 'variant_fitness'
 return 'raw_tabular_schema_pending'


def scan_csv(stream,aid,name,db):
 import numpy as np
 h=hashlib.sha256();expanded=0
 def lines():
  nonlocal expanded
  first=True
  while True:
   b=stream.readline(MAX_LINE+1)
   if not b:break
   if len(b)>MAX_LINE:raise ValueError('text line budget exceeded')
   expanded+=len(b)
   if expanded>MAX_EXPANDED:raise ValueError('expanded text budget exceeded')
   h.update(b);text=b.decode('utf-8-sig' if first else 'utf-8');first=False;yield text
 reader=csv.DictReader(lines(),delimiter='\t' if name.lower().endswith('.tsv') else ',')
 fields=reader.fieldnames or [];kind=schema_kind(set(fields));issues=collections.Counter();first=[];labels={};n=0;lengths=collections.Counter();usage=collections.Counter();probes=collections.Counter();cache={};identity_batch=[]
 seqcol='sequence' if 'sequence' in fields else next((f for f in fields if f.lower() in ['glms_variant_sequence','rna_sequence','genotype_sequence','seq','utr','utr_sequence','5utr','3utr']),None)
 numeric=[f for f in fields if re.search(r'(?i)kobs|kcat|^KM$|fitness|DMS_score|^label$|^target$|^reads$|signal_to_noise|half_life|^beta$|^rl$|^mrl$|^rss$|^te$|^expression$|^value$',f)]
 if kind=='long_nucleotide_probing':numeric=[f for f in ['reactivity_DMS_MaP','reactivity_2A3_MaP'] if f in fields]
 react=[f for f in fields if re.fullmatch('reactivity_\\d+',f)]
 error=[f for f in fields if re.fullmatch('reactivity_error_\\d+',f)]
 def fail(reason):
  issues[reason]+=1
  if len(first)<20:first.append({'row':n,'reason':reason})
 parts=Path(name).parts;strategy=parts[1] if len(parts)>2 and parts[0]=='CV_splits' else 'UNSPECIFIED'
 split='train' if 'train' in parts else 'test' if 'test' in parts else 'UNSPECIFIED'
 alias=re.sub(r'_(train|test)$','',Path(name).stem)
 for n,row in enumerate(reader,1):
  seq=row.get(seqcol) if seqcol else None
  if seqcol:
   if not isinstance(seq,str) or not seq:fail('EMPTY_SEQUENCE');seq=seq or ''
   cache_key=row.get('sequence_id') if kind=='long_nucleotide_probing' else None
   if cache_key is None or cache.get(cache_key)!=seq:
    if not set(seq)<=set('ACGTUN'):fail('NONCANONICAL_SEQUENCE')
    if cache_key is not None:
     if cache_key in cache:fail('SEQUENCE_ID_CONFLICT')
     if len(cache)>=2000000:raise ValueError('sequence-ID cache budget exceeded')
     cache[cache_key]=seq
   lengths[len(seq)]+=1
  usage[str(row.get('Usage') or 'UNSPECIFIED')]+=1
  if 'experiment_type' in row:probes[str(row['experiment_type'])]+=1
  for col in numeric:
   value=row.get(col);stats=labels.setdefault(col,Numeric())
   if value in (None,'','nan','NaN','NA','N/A'):stats.missing+=1
   else:
    try:stats.add(float(value),'continuous')
    except (TypeError,ValueError):stats.invalid+=1
  if kind=='long_nucleotide_probing':
   try:
    pos=int(row['position_id'])
    if not (0<=pos<=len(seq)):fail('POSITION_OUTSIDE_BOTH_0_AND_1_BASED_RANGES')
   except (ValueError,TypeError):fail('INVALID_POSITION_ID')
  if kind=='wide_nucleotide_probing':
   if len(react)<len(seq):fail('INSUFFICIENT_REACTIVITY_COLUMNS')
   for columns,label in [(react,'reactivity'),(error,'reactivity_error')]:
    if not columns:continue
    text=','.join(row[c] if row.get(c) not in (None,'') else 'nan' for c in columns)
    with __import__('warnings').catch_warnings():
     __import__('warnings').simplefilter('error')
     try:values=np.fromstring(text,sep=',')
     except (ValueError,Warning):fail('INVALID_VECTOR:'+label);continue
    if len(values)!=len(columns):fail('PARTIAL_VECTOR_PARSE:'+label);continue
    valid=np.isfinite(values);sentinel=valid&(values==-999)
    issues['MASKED_NONFINITE_'+label]+=int((~valid[:len(seq)]).sum())
    issues['RAW_MINUS999_'+label]+=int(sentinel[:len(seq)].sum())
    if len(values)>len(seq) and np.any(valid[len(seq):]&~sentinel[len(seq):]):fail('FINITE_MEASUREMENT_OUTSIDE_SEQUENCE:'+label)
  if kind=='variant_fitness':
   key=hashlib.sha256(json.dumps(['RNA-GYM:raw-sequence:v1',seq],separators=(',',':')).encode()).digest()
   label=hashlib.sha256(str(row.get('DMS_score')).encode()).digest()
   obs=hashlib.sha256(json.dumps(list(row.items()),ensure_ascii=False,separators=(',',':')).encode()).digest()
   identity_batch.append((aid,name,strategy,alias,split,n,key,label,obs))
   if len(identity_batch)>=10000:
    with db:db.executemany('INSERT INTO variants VALUES (?,?,?,?,?,?,?,?,?)',identity_batch)
    identity_batch=[]
 if identity_batch:
  with db:db.executemany('INSERT INTO variants VALUES (?,?,?,?,?,?,?,?,?)',identity_batch)
 return {'member':name,'format':kind,'actual_rows':n,'member_sha256':h.hexdigest(),'expanded_bytes_read':expanded,
   'schema_columns':fields,'sequence_column':seqcol,'sequence_length_histogram':dict(lengths),
   'numeric_columns_separate':{k:v.result() for k,v in labels.items()},'issues_or_qc_mask_counts':dict(issues),
   'first_issues':first,'usage_distribution':dict(usage),'probe_distribution':dict(probes),
   'original_strategy':strategy,'original_split':split,'dataset_alias':alias,
   'position_index_origin_verified':False if kind=='long_nucleotide_probing' else None,
   'label_semantics_verified':False,'units_verified':False,'rows_removed':0,
   'competition_label_status':'HOLD_COMPETITION_LABEL_TERMS' if 'solution_ribonanza' in name or any('private' in k.lower() for k in usage) else 'NOT_REVIEWED',
   'sequence_transformation':'NONE; FASTA/CSV representation rules only; no T/U substitution',
   'admission_gate_pass':False}


def scan_excel(stream,name):
 from openpyxl import load_workbook
 data=stream.read(128*1024**2+1)
 if len(data)>128*1024**2:raise ValueError('nested Excel compressed cap exceeded')
 with zipfile.ZipFile(io.BytesIO(data)) as z:
  if sum(i.file_size for i in z.infolist())>512*1024**2:raise ValueError('nested Excel expansion cap exceeded')
  for i in z.infolist():safe_member(i.filename,Path('/tmp/rnagym-archive-validation'))
 wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True,keep_links=False);sheets=[]
 try:
  for ws in wb.worksheets:
   nonempty=0;nonnumeric=0;numeric=0;nonfinite=0
   for row in ws.iter_rows(values_only=True):
    if any(v is not None for v in row):nonempty+=1
    for value in row:
     if isinstance(value,(int,float)):
      numeric+=1;nonfinite+=int(not math.isfinite(value))
     elif value is not None:nonnumeric+=1
   sheets.append({'sheet':ws.title,'nonempty_rows_including_headers':nonempty,'numeric_cells':numeric,'nonnumeric_cells':nonnumeric,'nonfinite_numeric_cells':nonfinite})
 finally:wb.close()
 return {'member':name,'format':'xlsx','state':'HOLD_ASSAY_COLUMN_SCHEMA','sheets':sheets,'formulas_executed':False,'external_links_loaded':False,'row_count_scope':'worksheet rows including headers, not experimental sample N','admission_gate_pass':False}


def fasta(stream,name):
 records=0;lengths=collections.Counter();chars=collections.Counter();seq=[]
 def finish():
  nonlocal records
  if seq:
   text=''.join(seq);records+=1;lengths[len(text)]+=1;chars.update(text);seq.clear()
 for b in stream:
  line=b.decode('utf-8').strip()
  if line.startswith('>'):finish()
  elif line:seq.append(line)
 finish()
 return {'member':name,'format':'alignment_or_fasta','records':records,'length_histogram':dict(lengths),'alphabet_counts':dict(chars),'canonicalization':'concatenate FASTA lines only; case, gaps and T/U preserved','admission_gate_pass':False}


def inspect_archive(raw,row,run,db):
 output=run/'reports/archive_qc'/(row['asset_id']+'.json');output.parent.mkdir(parents=True,exist_ok=True)
 if output.exists():return json.loads(output.read_text())
 members=[];scans=[];holds=[];expanded=0
 if zipfile.is_zipfile(raw):
  archive=zipfile.ZipFile(raw);items=archive.infolist()
  if len(items)>100000:raise ValueError('archive member count cap')
  for m in items:
   safe_member(m.filename,Path('/tmp/rnagym-archive-validation'))
   kind=stat.S_IFMT(m.external_attr>>16)
   if kind not in (0,stat.S_IFREG,stat.S_IFDIR) or m.flag_bits&1:raise ValueError('archive links/special/encrypted entries refused')
  expanded=sum(m.file_size for m in items)
  if expanded>MAX_EXPANDED:raise ValueError('archive expansion exceeds32GiB cap')
  iterable=[(m.filename,m.file_size,m.is_dir(),lambda item=m:archive.open(item)) for m in items]
 else:
  archive=tarfile.open(raw,'r:*');items=archive.getmembers()
  if len(items)>100000:raise ValueError('archive member count cap')
  for m in items:
   if m.isdir() and m.name in ('.','./'):continue
   safe_member(m.name,Path('/tmp/rnagym-archive-validation'))
   if not (m.isfile() or m.isdir()):raise ValueError('archive links/special entries refused')
  expanded=sum(m.size for m in items)
  if expanded>MAX_EXPANDED:raise ValueError('archive expansion exceeds32GiB cap')
  iterable=[(m.name,m.size,m.isdir(),lambda item=m:archive.extractfile(item)) for m in items]
 try:
  for name,size,isdir,opener in iterable:
   if isdir:continue
   members.append({'path':name,'declared_bytes':size})
   if any(s in name for s in ['.DS_Store','__MACOSX','.ipynb_checkpoints','__pycache__']) or name.endswith(('.py','.pyc','.sh','.md')):
    holds.append({'member':name,'state':'EXCLUDED_AUXILIARY_COPY_OR_OS_METADATA'});continue
   print(json.dumps({'asset_id':row['asset_id'],'member':name,'stage':'QC_STREAMING','declared_bytes':size}),flush=True)
   try:
    if name.lower().endswith(('.csv','.tsv')):
     with opener() as stream:scans.append(scan_csv(stream,row['asset_id'],name,db))
    elif name.lower().endswith('.xlsx'):
     with opener() as stream:scans.append(scan_excel(stream,name))
    elif name.lower().endswith(('.fa','.fasta','.afa','.a2m')):
     with opener() as stream:scans.append(fasta(stream,name))
    else:holds.append({'member':name,'state':'HOLD_FORMAT_SCHEMA','format':Path(name).suffix,'reason':'Not executed/unpickled; modality-specific parser or extraction required.'})
   except Exception as e:holds.append({'member':name,'state':'HOLD_QC','error':type(e).__name__+': '+str(e)[:300]})
 finally:archive.close()
 report={'asset_id':row['asset_id'],'source_sha256':sha256(raw),'state':'ARCHIVE_SCANNED_REVIEW_PENDING',
  'member_count':len(members),'declared_expanded_bytes':expanded,'members':members,'scanned_members':scans,'held_members':holds,
  'raw_file_unchanged':True,'rows_removed':0,'admission_gate_pass':False,'admitted':False,
  'not_checked':['full original assay identity and rights','all label units/replicate/negative-generation semantics','unparsed member formats','full biological/homology leakage and human review']}
 immutable_json(output,report);return report


def main():
 p=argparse.ArgumentParser(description=__doc__);p.add_argument('--run',type=Path,required=True);p.add_argument('--manifest',action='append',type=Path,required=True);p.add_argument('--identity-tag',default='main');a=p.parse_args();run=a.run.resolve()
 if not re.fullmatch('[A-Za-z0-9_-]+',a.identity_tag):raise ValueError('unsafe identity tag')
 db=sqlite3.connect(run/'reports'/('variant_split_identity-'+a.identity_tag+'.sqlite'),timeout=60);db.execute('PRAGMA cache_size=-65536')
 db.execute('CREATE TABLE IF NOT EXISTS variants(asset_id TEXT,member TEXT,strategy TEXT,source_alias TEXT,original_split TEXT,row_number INTEGER,input_hash BLOB,label_hash BLOB,observation_hash BLOB)')
 rows={}
 for f in a.manifest:rows.update(read_manifest(f))
 results=[]
 with StateDB(ROOT/'run-state.sqlite') as state:
  for row in rows.values():
   raw=get_target(row,ROOT/'raw_cache')
   if not raw.is_file() or not (raw.parent/'provenance.json').is_file():continue
   if not zipfile.is_zipfile(raw) and not str(raw).endswith(('.tar.gz','.tar.xz')):continue
   if sha256(raw)!=raw.parent.name:raise ValueError('raw hash mismatch')
   output=run/'reports/archive_qc'/(row['asset_id']+'.json')
   if not output.exists():
    with db:db.execute('DELETE FROM variants WHERE asset_id=?',(row['asset_id'],))
   try:report=inspect_archive(raw,row,run,db)
   except Exception as e:
    report={'asset_id':row['asset_id'],'state':'HOLD_ARCHIVE_QC','error':type(e).__name__+': '+str(e),'admission_gate_pass':False}
    immutable_json(run/'reports'/('archive-failed-'+row['asset_id']+'.json'),report)
   key=state.ensure(row);old=state.job(key)['stage']
   if old=='HASHED':state.transition(key,'PARSED')
   if report['state']=='HOLD_ARCHIVE_QC':state.transition(key,'HOLD_QC',detail={'error':report['error']})
   results.append({k:report.get(k) for k in ['asset_id','state','member_count','declared_expanded_bytes']});print(json.dumps(results[-1]),flush=True)
 db.execute('CREATE INDEX IF NOT EXISTS variant_split_key ON variants(source_alias,strategy,input_hash)');db.commit()
 overlap=[]
 for alias,strategy,key in db.execute('''SELECT source_alias,strategy,input_hash FROM variants WHERE original_split IN ('train','test') GROUP BY source_alias,strategy,input_hash HAVING count(DISTINCT original_split)>1'''):
  overlap.append({'source_alias':alias,'original_strategy':strategy,'input_hash':key.hex(),'status':'EXACT_INPUT_SHARED_ACROSS_OFFICIAL_SPLITS','rows_removed':0})
 immutable_json(run/'reports'/('official_variant_split_overlap-'+a.identity_tag+'.json'),{'exact_cross_split_keys':overlap,'admission_gate_pass':False,'seed_or_splits_rebuilt':False})
 immutable_json(run/'reports'/('archive_qc_summary-'+a.identity_tag+'.json'),{'results':results,'rows_removed':0,'admitted':0});db.close()


if __name__=='__main__':main()
